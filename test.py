import streamlit as st
from docx import Document
from fpdf import FPDF
import tempfile

def extract_headings(doc):
    headings = []
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith('Heading' or 'Heading1' or 'Heading2' or 'Heading3'):
            headings.append(paragraph.text)
    return headings

def create_pdf(headings):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for heading in headings:
        pdf.set_font("Arial", 'B', size=14)
        pdf.cell(200, 10, txt=heading, ln=True)
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, txt="\n\n\n")  # Space for writing
    return pdf

def main():
    st.title("SOW Automation GenAi")
    uploaded_file = st.file_uploader("Upload a Word document", type="docx")
    
    if uploaded_file is not None:
        with tempfile.NamedTemporaryFile(delete=False) as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = tmp_file.name
        
        doc = Document(tmp_file_path)
        headings = extract_headings(doc)
        
        if headings:
            pdf = create_pdf(headings)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_pdf:
                pdf.output(tmp_pdf.name)
                tmp_pdf_path = tmp_pdf.name
            
            with open(tmp_pdf_path, "rb") as f:
                st.download_button(
                    label="Download PDF",
                    data=f,
                    file_name="Extracted_Template.pdf",
                    mime="application/pdf"
                )
        else:
            st.write("No headings found in the document.")

if __name__ == "__main__":
    main()
